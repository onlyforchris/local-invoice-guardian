# -*- coding: utf-8 -*-
"""核心查重与分类的最小回归检查。"""
import io
import urllib.error

import app
import update


def keys(fields, md5="same"):
    return {key for _level, key in app.fingerprint(fields, md5)}


def main():
    a = {"no": "12345678", "amount_cents": 1000}
    b = {"no": "12345678", "amount_cents": 1100}
    assert keys(a) & keys(b), "同发票号码、金额识别不一致时必须提示"
    assert "md5|same" in keys(a), "有结构化字段时也必须保留文件内容指纹"
    assert app.is_within(r"C:\used\a.pdf", r"C:\used")
    assert not app.is_within(r"C:\used_backup\a.pdf", r"C:\used")
    assert app.classify_text("项目名称：打印纸")[0] == "office"
    assert app.classify_text("项目名称：酒店住宿")[0] == "travel"
    assert app.classify_text("项目名称：停车费")[0] == "transport"
    assert app.classify_text("高德打车网约车服务")[0] == "transport"
    assert app.classify_text("员工下午茶餐饮")[0] == "benefit"
    assert app.classify_text("生产生活服务餐费")[0] == "benefit"
    assert app.classify_text("项目名称：休闲零食食品")[0] == "benefit"
    assert app.classify_text("项目名称：宴请礼品")[0] == "hospitality"
    assert app.classify_text("顺丰收派服务")[0] == "office"
    assert app.classify_text("网约车客运服务费")[0] == "transport"
    assert app.classify_text("92号车用乙醇")[0] == "transport"
    assert app.classify_text("人工智能技术服务")[0] == "service"
    assert app.classify_text("熟肉制品酱板鸭")[0] == "benefit"
    broken = "发票号码：\x002\x006\x003\x001\x007\x009\x000\x007\x001\x005\x000\x009\x000\x000\x000\x007\x002\x000\x003\x003\n2026^t09g\b07\n价税合计（小写）\x00¥\x001\x00.\x009\x006"
    fields = app.extract_fields(broken)
    assert fields["no"] == "26317907150900072033"
    assert fields["date"] == "2026-09-07"
    assert fields["amount_cents"] == 196
    # 数电票新版式：标签与值分离 + 数值逐位拆空格
    sparse = ("电子发票（普通发票） 发票号码：\n开票日期：\n购\n买\n方\n信\n息\n"
              "2 6 3 3 2 0 0 0 0 0 7 7 7 6 3 8 3 5 5 1\n"
              "2 0 2 6 年0 9 月0 9 日\n"
              "上海亿流科技有限公司\n9131011259474644XJ\n"
              "金匠匠（杭 州）餐饮管理有限公司\n91330106MADGKX7Q15\n"
              "¥5 8 . 0 0 ¥0 . 5 8\n伍拾捌圆伍角捌分 ¥5 8 . 5 8\n"
              "* 生产生活服务* 餐饮服务 1 %5 8 . 0 0 0 . 5 8")
    f2 = app.extract_fields(sparse, company_names=["上海亿流科技有限公司"])
    assert f2["no"] == "26332000007776383551", f2["no"]
    assert f2["date"] == "2026-09-09", f2["date"]
    assert f2["amount_cents"] == 5858, f2["amount_cents"]
    assert f2["buyer"] == "上海亿流科技有限公司"
    # 注：NFKC 归一化会把全角括号统一为半角，属于预期行为（利于查重比对一致性）
    assert f2["seller"] == "金匠匠(杭州)餐饮管理有限公司", f2["seller"]
    assert f2["kind"] == "数电票"
    assert not app.validate_fields(f2), app.validate_fields(f2)
    assert app.extract_fields("价税合计（大写）：壹佰圆肆角伍分 （小写）：100.45")["amount_cents"] == 10045
    party = app.extract_fields(
        "名称: 上海亿流科技有限公司\n统一社会信用代码: 9131011259474644XJ\n"
        "名称: 杭州市西湖区国娣餐饮店\n统一社会信用代码: 92330106MA2GKHP698",
        company_names=["上海亿流科技有限公司"])
    assert party["seller"] == "杭州市西湖区国娣餐饮店"
    # 个体小商户后缀（小吃店/商行等）也必须能识别
    assert "上海市长宁区杨震小吃店" in app._companies(
        "9131011259474644XJ\n上海市长宁区杨震小吃店\n92310105MAC4EUBHX7")
    # 杂串粘连（数电票机器可读区）：b017 不得被当成公司名前缀
    dirty = "01,32,,26317200000008021556,16.40,20260908,,b017上海亿流科技有限公司\n" \
            "9131011259474644XJ\n上海市长宁区杨震小吃店\n92310105MAC4EUBHX7"
    dcomps = app._companies(dirty)
    assert not any("b017" in c for c in dcomps), dcomps
    df = app.extract_fields(dirty, company_names=["上海亿流科技有限公司"])
    assert df["buyer"] == "上海亿流科技有限公司", df["buyer"]
    assert df["seller"] == "上海市长宁区杨震小吃店", df["seller"]
    joined = app.extract_fields(
        "26332000007717763206\n2026年09月07日\n上海亿流科技有限公司\n"
        "9131011259474644XJ杭州市西湖区显光烧饼店\n92330106MA2BK14X8G",
        company_names=["上海亿流科技有限公司"])
    assert joined["seller"] == "杭州市西湖区显光烧饼店"
    compact = "上海亿流科技有限公司 上海⾦拱⻔⻝品有限公司\n9131011259474644XJ 913100006072048513"
    fields = app.extract_fields(compact, company_names=["上海亿流科技有限公司"])
    assert fields["seller"] == "上海金拱门食品有限公司"
    assert app.validate_categories(app.DEFAULT_CATEGORIES)
    assert update.version_key("v1.1.10") > update.version_key("1.1.9")
    assert "无效" in app.friendly_ocr_error(urllib.error.HTTPError("", 401, "", {}, None))
    assert "频率" in app.friendly_ocr_error(urllib.error.HTTPError("", 429, "", {}, None))

    # xlsx 分类汇总导出
    from openpyxl import load_workbook
    recs = [
        {"path": "a", "fname": "a.pdf", "folder": "F", "amount_cents": 1470,
         "date": "2026-08-02", "cat_label": "交通费", "seller": "s1",
         "is_used": False, "dups": []},
        {"path": "b", "fname": "b.pdf", "folder": "F", "amount_cents": 200,
         "date": "2026-08-01", "cat_label": "交通费", "seller": "s2",
         "is_used": False, "dups": []},
        {"path": "c", "fname": "c.pdf", "folder": "F", "amount_cents": 77700,
         "date": "2026-08-03", "cat_label": "差旅费", "seller": "s3",
         "is_used": True, "dups": []},
        {"path": "d", "fname": "d.pdf", "folder": "F", "amount_cents": None,
         "date": None, "cat_label": "交通费", "seller": "s4",
         "is_used": False, "dups": []},
    ]
    buf = app.export_xlsx(recs)
    wb = load_workbook(filename=io.BytesIO(buf))
    ws = wb["分类汇总"]
    assert ws.cell(1, 1).value == "差旅费", "列顺序须按分类目录（差旅费在交通费前）"
    assert ws.cell(1, 2).value == "交通费"
    assert ws.cell(2, 2).value == 2, "金额按日期升序：0.8-1 的 2 元在前"
    assert ws.cell(3, 2).value == 14.7
    assert ws.cell(2, 1).value == 777
    assert ws.cell(5, 1).value == "=SUM(A2:A3)", "合计行统一覆盖数据区（空单元格按0计）"
    assert ws.cell(5, 2).value == "=SUM(B2:B3)"
    assert ws.cell(1, 3).value == "总计"
    ws2 = wb["明细"]
    assert ws2.cell(2, 4).value == "s2"
    assert app.record_in_scope({"is_used": False, "dups": [{"is_used": True}]}, "reused")
    assert not app.record_in_scope({"is_used": True, "dups": [{"is_used": True}]}, "reused")
    assert app.record_in_scope({"is_used": True, "in_watch": True, "dups": []}, "new")
    assert not app.record_in_scope({"is_used": False, "in_watch": False, "dups": []}, "new")
    assert not app.record_in_scope({"is_used": False, "dups": []}, "dup")

    # T2 字段自校验
    assert any("位数" in x for x in app.validate_fields({"no": "12345"}))
    assert not app.validate_fields({"no": "12345678"})
    assert not app.validate_fields({"no": "2" + "0" * 19})
    assert any("晚于今天" in x for x in app.validate_fields({"date": "2099-01-01"}))
    assert any("自洽" in x for x in app.validate_fields(
        {"amount_cents": 10000, "amount_excl_cents": 5000, "tax_cents": 300}))
    assert not app.validate_fields(
        {"amount_cents": 10000, "amount_excl_cents": 9434, "tax_cents": 566})

    # T5 抬头校验
    assert app.check_buyer({"buyer": "上海亿流科技有限公司"}, ["上海亿流"])[0] == "符合"
    assert app.check_buyer({"buyer": "杭州市西湖区国娣餐饮店"}, ["上海亿流"])[0] == "不符"
    assert app.check_buyer({"buyer": None}, [])[0] == "未校验"

    # N2 永久库：无 md5 时不得生成 exact 键，否则条目互相误判
    assert not any(k.startswith("md5|") for _w, k in app.fingerprint({"no": "1" * 20}, None))
    assert any(k.startswith("md5|") for _w, k in app.fingerprint({"no": "1" * 20}, "abc"))
    assert app.used_item_key({"no": "1" * 20, "amount_cents": 100}) == "no+amt|%s|100" % ("1" * 20)
    assert app.used_item_key({"no": None, "amount_cents": None}) == ""

    # N3 历史台账导入
    hist = app.parse_history_csv("发票号码,金额,开票日期\n24312000000112345678,100.45,2026-08-01\n")
    assert len(hist) == 1 and hist[0]["no"] == "24312000000112345678" and hist[0]["amount_cents"] == 10045

    # N4 归档规范命名
    nm = app.archive_name({"date": "2026-08-01", "seller": "杭州*餐饮/店", "amount_cents": 10045,
                           "no": "24312000000112345678", "fname": "a.pdf"})
    assert nm.startswith("2026-08-01_杭州餐饮店_100.45_24312000000112345678") and nm.endswith(".pdf")

    # OFD（数电票官方格式）：zip + XML 解析，零第三方依赖
    import os
    import shutil
    import tempfile
    import zipfile
    tmpd = tempfile.mkdtemp(prefix="ofd-test-")
    try:
        ofd_p = os.path.join(tmpd, "t.ofd")
        with zipfile.ZipFile(ofd_p, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("OFD.xml",
                       '<?xml version="1.0" encoding="UTF-8"?>'
                       '<ofd:OFD xmlns:ofd="http://www.ofdspec.org/2016">'
                       '<ofd:DocBody><ofd:DocRoot>Doc_0/Document.xml</ofd:DocRoot></ofd:DocBody></ofd:OFD>')
            z.writestr("Doc_0/Document.xml",
                       '<?xml version="1.0" encoding="UTF-8"?>'
                       '<ofd:Document xmlns:ofd="http://www.ofdspec.org/2016"><ofd:Pages>'
                       '<ofd:Page ID="1" BaseLoc="Pages/Page_0/Content.xml"/></ofd:Pages></ofd:Document>')
            z.writestr("Doc_0/Pages/Page_0/Content.xml",
                       '<?xml version="1.0" encoding="UTF-8"?>'
                       '<ofd:Content xmlns:ofd="http://www.ofdspec.org/2016"><ofd:Layer ID="1">'
                       '<ofd:TextObject ID="1" Boundary="10 10 100 8"><ofd:TextCode>电子发票（普通发票）</ofd:TextCode></ofd:TextObject>'
                       '<ofd:TextObject ID="2" Boundary="10 20 100 8"><ofd:TextCode>发票号码：</ofd:TextCode></ofd:TextObject>'
                       '<ofd:TextObject ID="3" Boundary="60 20 100 8"><ofd:TextCode>24312000000112345678</ofd:TextCode></ofd:TextObject>'
                       '<ofd:TextObject ID="4" Boundary="10 30 100 8"><ofd:TextCode>开票日期：2026年09月09日</ofd:TextCode></ofd:TextObject>'
                       '<ofd:TextObject ID="5" Boundary="10 40 100 8"><ofd:TextCode>上海亿流科技有限公司</ofd:TextCode></ofd:TextObject>'
                       '<ofd:TextObject ID="6" Boundary="10 50 100 8"><ofd:TextCode>金匠匠（杭州）餐饮管理有限公司</ofd:TextCode></ofd:TextObject>'
                       '<ofd:TextObject ID="7" Boundary="10 60 100 8"><ofd:TextCode>价税合计（小写）¥58.58</ofd:TextCode></ofd:TextObject>'
                       '</ofd:Layer></ofd:Content>')
        otxt = app.read_ofd_text(ofd_p)
        assert "24312000000112345678" in otxt, otxt
        of = app.extract_fields(otxt, company_names=["上海亿流科技有限公司"])
        assert of["no"] == "24312000000112345678"
        assert of["date"] == "2026-09-09"
        assert of["amount_cents"] == 5858
        assert of["kind"] == "数电票"
        assert of["seller"] == "金匠匠(杭州)餐饮管理有限公司"
        assert ".ofd" in app.INVOICE_EXTS
        assert app.read_ofd_text(os.path.join(tmpd, "nope.ofd")) == ""
    finally:
        shutil.rmtree(tmpd, ignore_errors=True)

    # 本地识别为可选依赖：未安装时 local_ocr_ready() 必须为 False 且不抛异常
    assert isinstance(app.local_ocr_ready(), bool)

    # 「数字在前、¥ 在后」版式（滴滴/高德等电子行程发票实测）：标签被逐字拆开、
    # 金额行形如 “合 计 366.67¥ 11.00¥”“（ 小 写 ） 377.67¥叁佰…”。
    # 曾经的 bug：稀疏归一化把两笔拼成 “366.67¥11.00¥”，¥ 兜底抓出假金额 11.00。
    rev = app.extract_fields(
        "开票日期 : 2026年08月17日\n发票号码 : 26317000002998211318\n"
        "名称： 上海亿流科技有限公司\n统一社会信用代码/纳税人识别号： 9131011259474644XJ\n"
        "名称： 上海滴滴畅行科技有限公司\n"
        "合 计 366.67¥ 11.00¥\n"
        "价 税 合 计 （ 大 写 ） （ 小 写 ） 377.67¥叁佰柒拾柒圆陆角柒分\n")
    assert rev["no"] == "26317000002998211318"
    assert rev["amount_cents"] == 37767, "¥ 后置版式必须以价税合计 377.67 为准，不得抓到假的 11.00"
    assert rev["amount_excl_cents"] == 36667 and rev["tax_cents"] == 1100
    assert rev["seller"] == "上海滴滴畅行科技有限公司"
    assert app.validate_fields(rev) == [], "377.67 = 366.67 + 11.00，自检应通过"

    # 审核底稿导出
    from openpyxl import load_workbook as _lw
    wrecs = [{"path": "a", "fname": "a.pdf", "folder": "张三", "in_watch": True,
              "amount_cents": 1000, "date": "2026-08-01", "seller": "s", "no": "n1",
              "cat_label": "交通费", "verify_status": "已查验通过", "check_issues": [],
              "reject_reason": "", "is_used": False},
             {"path": "b", "fname": "b.pdf", "folder": "李四", "in_watch": True,
              "amount_cents": 2000, "date": "2026-08-02", "seller": "s", "no": "n2",
              "cat_label": "差旅费", "verify_status": "查验异常", "check_issues": ["号码位数异常"],
              "reject_reason": "疑似假票", "is_used": False}]
    wb3 = _lw(filename=io.BytesIO(app.export_worksheet(wrecs)))
    ws3 = wb3["审核底稿"]
    assert ws3.cell(1, 9).value == "查验状态"
    assert ws3.cell(2, 3).value == "a.pdf" and ws3.cell(3, 3).value == "b.pdf"
    assert wb3["退回清单"].cell(2, 3).value == "n2", "退回清单只含有退回原因的票"
    print("core checks: ok")


if __name__ == "__main__":
    main()
