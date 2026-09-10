# -*- coding: utf-8 -*-
"""
发票管家 Invoice Manager —— 本地发票内容级查重 + 报销分类桌面工具
================================================================
能力：
  1. 读取本地可配置文件夹（收集箱 / 已使用批次文件夹 / 归档目录）内的发票
  2. 解析发票正文：文本型 PDF 用 pypdf 直接提字；无文本层 PDF / jpg / png 走
     智谱 GLM-4V OCR（需 ZHIPUAI_API_KEY，未配置时自动降级并提示）
  3. 内容级查重：不依赖文件名。指纹 = (发票号码+金额) / (发票号码) / (销方+金额+日期) / 文件MD5
  4. 报销分类：13 类费用科目（关键词规则，发票正文+文件名仅作分类线索），可人工改分类并记忆
  5. 台账持久化(JSON)、导出 CSV、把发票复制提取到归档目录

用法：
  python app.py                 # 启动本地服务(127.0.0.1:8765)并自动开浏览器
  python app.py --port 9000 --no-browser
  python app.py --test <dir> [dir...]   # 命令行只读自测
"""
import argparse
import base64
import glob
import hashlib
import io
import json
import os
import re
import shutil
import sys
import tempfile
import threading
import time
import unicodedata
import urllib.error
import urllib.request
import webbrowser
from decimal import Decimal
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs

APP_DIR = os.path.dirname(os.path.abspath(__file__))
WORKBUDDY_PACKAGES = os.path.expanduser(r"~/.workbuddy/binaries/python/envs/site-packages")
if os.path.isdir(WORKBUDDY_PACKAGES):
    # 仅作兜底，避免覆盖当前 Python 中与其 ABI 匹配的 Pillow 等二进制包。
    sys.path.append(WORKBUDDY_PACKAGES)
try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None
try:
    import pypdfium2 as pdfium
    from PIL import Image
except Exception:
    pdfium = None
    Image = None
try:
    # 本地 OCR 为可选依赖：未安装时自动降级到 AI 识别，不影响主流程
    from rapidocr_onnxruntime import RapidOCR
except Exception:
    RapidOCR = None

LEDGER = os.path.join(APP_DIR, "invoice_ledger.json")
# 已报销永久库：脱离发票文件位置存在，即使原文件被移动/删除也能持续查重
USED_LEDGER = os.path.join(APP_DIR, "used_ledger.json")
BACKUP_DIR = os.path.join(APP_DIR, "backups")
CONFIG = os.path.join(APP_DIR, "config.json")
STATIC = os.path.join(APP_DIR, "index.html")
ENGINE_VER = 10  # 引擎版本；升级后旧台账自动失效重解析（改解析逻辑必须同步升级！）
APP_VERSION = "1.4.3"
INVOICE_EXTS = {".pdf", ".ofd", ".jpg", ".jpeg", ".png", ".bmp", ".webp"}
BUYER_DEFAULT = []
BACKUP_KEEP = 5
# 台账健康状态：损坏时不再静默清空，必须让用户看见并可恢复
LEDGER_STATUS = {"error": None, "backups": [], "record_count": 0}
VERIFY_CHOICES = ("未查验", "已查验通过", "查验异常")
REJECT_CHOICES = ("", "重复报销", "抬头/税号不符", "疑似假票", "金额不符", "票面信息不全", "其他")

# ---------- 报销分类科目（13 类）----------
CATALOG = [
    {"id": "office", "label": "办公费", "note": "办公用品、打印纸、快递、饮用水、绿植"},
    {"id": "travel", "label": "差旅费", "note": "机票、火车票、住宿、打车、订票手续费"},
    {"id": "transport", "label": "交通费", "note": "网约车、加油、停车、过路、租车"},
    {"id": "communication", "label": "通讯费", "note": "办公电话、宽带、通信服务"},
    {"id": "hospitality", "label": "业务招待费", "note": "餐饮招待、茶叶、礼品、宴请"},
    {"id": "meeting", "label": "会议费", "note": "场地、资料"},
    {"id": "training", "label": "培训费", "note": "培训、报名费"},
    {"id": "advertising", "label": "广告宣传费", "note": "广告服务费、制作费、推广费"},
    {"id": "service", "label": "服务费", "note": "技术服务费、招聘费"},
    {"id": "rd", "label": "研发费用", "note": "专利/软著、研发设备、认证检测"},
    {"id": "property", "label": "房租物业水电", "note": "租赁、物业、水电"},
    {"id": "benefit", "label": "人事福利", "note": "团建、下午茶、餐饮、零食、食品"},
    {"id": "other", "label": "其他/待分类", "note": "建议人工补充规则"},
]
CAT_IDS = [c["id"] for c in CATALOG]

# 分类关键词规则（有序，先命中更具体；文本=文件名+销方+品名+正文片段）
RULES = [
    ("office", ["办公用品", "打印纸", "快递", "收派服务", "物流网络", "饮用水", "绿植", "文具", "耗材", "硒鼓", "墨盒", "a4纸"]),
    ("travel", ["机票", "火车票", "住宿", "打车", "订票手续费", "行程单", "高铁", "动车", "酒店", "宾馆"]),
    ("transport", ["网约车", "客运服务费", "加油", "车用乙醇", "汽油", "柴油", "停车", "过路", "租车", "滴滴", "出租车", "通行费", "代驾"]),
    ("communication", ["办公电话", "宽带", "通信服务", "中国移动", "中国联通", "中国电信", "话费"]),
    ("hospitality", ["餐饮招待", "宴请", "茶叶", "礼品"]),
    ("meeting", ["会议场地", "会议资料", "会议费"]),
    ("training", ["培训", "报名费", "课程", "训练营", "考试费"]),
    ("advertising", ["广告服务费", "制作费", "推广费", "广告宣传"]),
    ("service", ["技术服务费", "技术服务", "招聘费", "软件服务", "云服务", "咨询服务"]),
    ("rd", ["专利", "软著", "软件著作权", "研发设备", "认证检测", "认证费", "检测费"]),
    ("property", ["房屋租赁", "租赁费", "物业", "水费", "电费", "水电"]),
    ("benefit", ["团建", "下午茶", "员工餐", "餐饮", "餐费", "外卖", "咖啡", "奶茶", "零食", "食品", "熟肉制品", "酱板鸭", "水果", "牛奶", "乳制品", "纯奶"]),
]

DEFAULT_CATEGORIES = [
    {**c, "keywords": next((list(kws) for cid, kws in RULES if cid == c["id"]), [])}
    for c in CATALOG
]
STATE_LOCK = threading.RLock()

COMPANY_RE = re.compile(
    # 公司名必须以汉字/全角括号开头：防止杂串粘连（如“b017上海亿流…”被当成公司名）
    r"(?=[\u4e00-\u9fa5（(])"
    r"[\u4e00-\u9fa5A-Za-z0-9（）()·]+?"
    r"(?:有限公司|有限责任公司|股份有限公司|个体工商户|合伙企业|工作室"
    r"|服务中心|购物中心|合作社|事务所|招待所|门市部|经营部|加工厂|汽修厂|幼儿园|食堂"
    r"|商行|超市|餐厅|饭店|酒店|宾馆|大药房|药店|诊所|美容院|医院"
    r"|便利店|小吃店|水果店|奶茶店|咖啡店|甜品店|烘焙店|理发店|打印店|图文店"
    r"|店)")


def get_categories(cfg=None):
    cfg = load_config() if cfg is None else cfg
    cats = cfg.get("categories")
    return cats if isinstance(cats, list) and cats else DEFAULT_CATEGORIES


def classify_text(text, categories=None):
    t = (text or "").lower()
    t = re.sub(r"\s+", "", t)
    if not t:
        return "other", "无内容"
    matches = []
    for order, cat in enumerate(categories or DEFAULT_CATEGORIES):
        for k in cat.get("keywords", []):
            if k in t:
                matches.append((len(k), -order, cat["id"], k))
    if matches:
        _length, _order, cid, keyword = max(matches)
        return cid, keyword
    return "other", "未命中规则"


# ---------- 字段归一化 ----------
def _strip_ws(s):
    return re.sub(r"\s+", "", s or "")


def norm_amount_cents(raw):
    if raw is None:
        return None
    try:
        return int(Decimal(re.sub(r"[^\d.]", "", str(raw))) * 100)
    except Exception:
        return None


def _valid_date(y, m, d):
    try:
        y, m, d = int(y), int(m), int(d)
        return 1900 <= y <= 2100 and 1 <= m <= 12 and 1 <= d <= 31
    except Exception:
        return False


def norm_date(raw):
    if not raw:
        return None
    m = re.search(r"(\d{4})[年\-/.](\d{1,2})[月\-/.](\d{1,2})", str(raw))
    if m and _valid_date(*m.groups()):
        return "%04d-%02d-%02d" % tuple(int(x) for x in m.groups())
    m = re.search(r"(\d{4})(\d{2})(\d{2})", str(raw))
    if m and _valid_date(*m.groups()):
        return "%04d-%02d-%02d" % tuple(int(x) for x in m.groups())
    return None


# ---------- 从正文提取字段（PDF 文本与 OCR 文本共用）----------
def _companies(text):
    out = []
    # NFKC 会把全角括号转成半角，而公司名中可能夹有空格（如“杭 州”），
    # 字符类不含空白会从空格处截断。补跑一份去空白文本，找回完整公司名。
    # 源2：仅压缩汉字/括号之间的空格与制表符（不含换行，避免跨行粘连），
    # 修复“杭 州”这类断裂；ASCII（税号）不受影响，避免税号与公司名粘连
    sources = [(text or "", True),
               (re.sub(r"(?<=[\u4e00-\u9fa5（）()·])[ \t\u00a0]+(?=[\u4e00-\u9fa5（）()·])",
                       "", text or ""), False)]
    for src, allow_taxgap in sources:
        candidates = COMPANY_RE.findall(src)
        candidates += re.findall(r"名称[:：]\s*([^\n]{2,60})", src)
        if allow_taxgap:  # “税号之间取文本”仅用于原文本；去空白文本会把税号吞进公司名
            taxes = list(re.finditer(r"(?<![0-9A-Z])[0-9A-Z]{15,20}(?![0-9A-Z])", src))
            for a, b in zip(taxes, taxes[1:]):
                # 间隙里再用公司名正则提取，避免把“金额,日期,杂串”整段吞成候选
                candidates += COMPANY_RE.findall(src[a.end():b.start()])
        for c in candidates:
            c = re.sub(r"\s+", "", c).strip(":：")
            if (4 <= len(re.findall(r"[\u4e00-\u9fa5]", c)) <= 40
                    and not re.search(r"发票|统一社会信用|纳税人识别号|项目名称", c)
                    and not norm_date(c)
                    and c not in out):
                out.append(c)
    # 截断候选（如“州)餐饮管理有限公司”）已被完整版本包含时丢弃，保留最完整公司名
    return [c for c in out if not any(c != o and c in o for o in out)]


def extract_fields(text, fname="", company_names=None):
    """返回字段 dict。金额以「价税合计（小写）」为准。"""
    # 一些通行费 PDF 会在每个字符前插入 NUL；先清理再做结构化匹配。
    t = unicodedata.normalize("NFKC", (text or "").replace("\x00", "")).translate(
        str.maketrans({"⻔": "门", "⻝": "食"}))
    # 数电票新版式：标签与值分离（“发票号码：\n”后无值），且数值被逐位拆开
    # （如 “2 6 3 3 ...”、“¥5 8 . 5 8”、“2 0 2 6 年0 9 月0 9 日”）。
    # 把「数字/¥/小数点 之间仅隔空白」的稀疏序列压缩还原，否则号码/日期/金额全部漏识别。
    t = re.sub(r"(?<=[0-9¥￥.])[ \t\r\n]+(?=[0-9¥￥.])", "", t)
    # 关键金额标签被逐字拆开（“合 计”“价 税 合 计”“（ 小 写 ）”），先归位再做金额匹配；
    # 否则标签锚点全部失配，金额只能靠 ¥ 兜底，容易抓错。
    t = re.sub(r"价\s*税\s*合\s*计", "价税合计", t)
    t = re.sub(r"合\s*计", "合计", t)
    t = re.sub(r"小\s*写", "小写", t)
    t2 = re.sub(r"(\d),(?=\d{3})", r"\1", t)
    f = {"no": None, "code": None, "date": None, "amount_cents": None,
         "seller": None, "buyer": None, "kind": "其他", "items": []}

    m = re.search(r"发票号码[:：\s]*([0-9]{8,20})", t2)
    if not m:
        m = re.search(r"([0-9]{20})", t2)
    f["no"] = m.group(1) if m else None
    m = re.search(r"发票代码[:：\s]*([0-9]{10,12})", t2)
    f["code"] = m.group(1) if m else None

    # 日期：优先“开票日期”标签，其次靠近发票号码/价税合计位置的独立日期
    m = re.search(r"开票日期[^0-9]{0,8}([0-9]{4}[年\-/.]\d{1,2}[月\-/.]\d{1,2})", t2)
    if not m:
        m = re.search(r"开票日期[^0-9]{0,8}(\d{8})", t2)
    if m:
        f["date"] = norm_date(m.group(1))
    if not f["date"]:
        anchor = t2.find(f["no"]) if f["no"] else t2.find("价税合计")
        if anchor < 0:
            anchor = 0
        dates = list(re.finditer(r"(\d{4})[年\-/.](\d{1,2})[月\-/.](\d{1,2})", t2))
        if dates:
            dates = sorted(dates, key=lambda x: abs(x.start() - anchor))
            f["date"] = norm_date(dates[0].group(0))
    if not f["date"] and f["no"]:
        # 兼容形如 2026^t09g\b07 的损坏文本层，仅在发票号后的小窗口内兜底。
        tail = t2[t2.find(f["no"]) + len(f["no"]):][:120]
        m = re.search(r"(20\d{2})\D{0,6}(\d{2})\D{0,6}(\d{2})", tail)
        if m and _valid_date(*m.groups()):
            f["date"] = "%04d-%02d-%02d" % tuple(int(x) for x in m.groups())

    # 金额：价税合计（小写）→ 价税合计后首个数字 → （小写）→ 合计 → 最大 ¥
    # 「数字在前、¥ 在后」版式（如 “…377.67¥叁佰…”）由第二条模式覆盖，
    # 否则这类票只能落到 ¥ 兜底，而稀疏归一化会把 “366.67¥ 11.00¥” 拼成
    # “366.67¥11.00¥” 造出假 “¥11.00”，金额直接抓错。
    amt = None
    for pat in (r"价税合计[^¥￥]{0,24}[¥￥]\s*([0-9]+\.\d{2})",
                r"价税合计[^0-9]{0,20}([0-9]+\.\d{2})",
                r"[（(]\s*小写\s*[)）]\s*[:：]?\s*[¥￥]?\s*([0-9]+\.\d{2})",
                r"(?:合计|合计金额)[^0-9¥￥]{0,12}([0-9]+\.\d{2})"):
        m = re.search(pat, t2)
        if m:
            amt = m.group(1)
            break
    if amt is None:
        # ¥ 兜底同时收「¥在数字前」和「¥在数字后」两种写法，取最大
        nums = re.findall(r"[¥￥]\s*([0-9]+\.\d{2})", t2) + \
               re.findall(r"([0-9]+\.\d{2})[ \t]*[¥￥]", t2)
        if nums:
            amt = max(nums, key=lambda x: Decimal(x))
    f["amount_cents"] = norm_amount_cents(amt)

    # 不含税合计与税额：仅在同行出现两个金额时才采信，用于金额自洽校验
    f["amount_excl_cents"] = None
    f["tax_cents"] = None
    m = re.search(r"合\s*计[^0-9¥￥\n]{0,10}[¥￥]?\s*([0-9]+\.\d{2})[^0-9¥￥\n]{0,10}[¥￥]?\s*([0-9]+\.\d{2})", t2)
    if m:
        f["amount_excl_cents"] = norm_amount_cents(m.group(1))
        f["tax_cents"] = norm_amount_cents(m.group(2))

    # 购销双方：公司名识别，购方=本公司（config 可改）
    buyers = company_names or BUYER_DEFAULT
    comps = _companies(t2)
    if comps:
        b = next((c for c in comps if any(bk in c for bk in buyers)), comps[0])
        s = next((c for c in comps if c != b), None)
        f["buyer"], f["seller"] = b, s

    # 类型：数电票标题为“电子发票（普通/专用发票）”（NFKC 后括号为半角），
    # 与老版“增值税普通/专用发票”区分，便于财务按票种核验
    if ("数电" in t2 or "电子发票（普通" in t2 or "电子发票(普通" in t2
            or "电子发票（专用" in t2 or "电子发票(专用" in t2):
        f["kind"] = "数电票"
    elif "增值税专用发票" in t2:
        f["kind"] = "增值税专用发票"
    elif "增值税普通发票" in t2:
        f["kind"] = "增值税普通发票"
    elif "行程单" in t2:
        f["kind"] = "行程单"
    elif "通行费" in t2:
        f["kind"] = "通行费发票"
    elif "出租车" in t2:
        f["kind"] = "出租车票"

    # 品名（项目名称表格首行明细）+ 服务编码（如 *交通运输服务*客运服务费）
    items = []
    m = re.search(r"项目名称[^\n]{0,12}\n", t2)
    if m:
        seg = t2[m.end():m.end() + 400]
        for line in seg.splitlines():
            line = line.strip()
            if not line or re.search(r"规格型号|单位|数量|单价|金额|税率|税额|价税合计|备注|合计", line):
                break
            if len(line) > 3:
                items.append(line[:80])
    f["items"] = items[:6]
    f["svc"] = re.findall(r"\*[^*\n]{1,30}\*([^*\n]{2,40})", t2)[:4]

    return f


def parse_history_csv(text):
    """解析财务已有的 Excel/CSV 台账（列顺序不限），最小集 = 号码 + 金额。"""
    import csv as _csv
    out = []
    for row in _csv.reader(io.StringIO(text or "")):
        cells = [c.strip() for c in row if c and c.strip()]
        if not cells:
            continue
        no = date = seller = None
        cents = None
        for c in cells:
            if date is None:
                d = norm_date(c)
                if d:
                    date = d
                    continue
            if no is None and re.fullmatch(r"\d{8}|\d{20}", c):
                no = c
                continue
            if cents is None and re.fullmatch(r"[¥￥]?\s*-?\d+\.\d{2}", c):
                cents = norm_amount_cents(c)
                continue
            if seller is None and len(c) <= 40 and re.search(r"[\u4e00-\u9fa5]{4}", c) \
                    and not re.search(r"已报|未报|报销|发票|状态|分类|合计", c):
                seller = c
        if no or (cents is not None and date):
            out.append({"no": no, "code": None, "amount_cents": cents, "date": date,
                        "seller": seller, "fname": "（历史台账导入）", "path": "",
                        "buyer": None, "kind": None, "cat_label": None})
    return out


def archive_name(r):
    """归档规范命名：日期_销售方_金额_号码，便于日后按文件名检索。

    号码缺失时退回原文件名，避免多张未识别票挤成同一个名字。
    """
    ext = os.path.splitext(r.get("fname") or "")[1]
    if not r.get("no"):
        stem = os.path.splitext(r.get("fname") or "未识别发票")[0]
        return re.sub(r'[\\/:*?"<>|\r\n\t]', "", stem)[:60] + ext
    safe = re.sub(r'[\\/:*?"<>|\r\n\t]', "", r.get("seller") or "未知销售方")[:20]
    return "_".join([r.get("date") or "无日期", safe,
                     "%.2f" % ((r.get("amount_cents") or 0) / 100),
                     r.get("no")]) + ext


export_worksheet_head = ["序号", "提交人/批次", "文件名", "发票号码", "开票日期",
                         "销售方", "金额(元)", "费用分类", "查验状态",
                         "合规校验", "退回原因", "报销状态"]


def export_worksheet(recs):
    """审核底稿：明细 + 按提交人小计 + 合计 + 退回清单，可直接打印/存档。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    rows = [r for r in recs if r.get("in_watch")]
    rows.sort(key=lambda r: (r.get("folder") or "", r.get("date") or "", r.get("fname") or ""))
    wb = Workbook()
    ws = wb.active
    ws.title = "审核底稿"
    bold = Font(bold=True)
    ws.append(export_worksheet_head)
    for c in ws[1]:
        c.font = bold
    for i, r in enumerate(rows, 1):
        issues = "；".join(r.get("check_issues") or [])
        ws.append([i, r.get("folder") or "", r.get("fname") or "", r.get("no") or "",
                   r.get("date") or "", r.get("seller") or "",
                   (r.get("amount_cents") or 0) / 100, r.get("cat_label") or "",
                   r.get("verify_status") or "未查验", issues,
                   r.get("reject_reason") or "",
                   "已报销" if r.get("is_used") else "待报销"])
    n = len(rows)
    # 按提交人小计
    start = n + 3
    ws.cell(row=start, column=1, value="按提交人/批次小计").font = bold
    sums = {}
    for r in rows:
        sums[r.get("folder") or "（未分组）"] = sums.get(r.get("folder") or "（未分组）", 0) \
            + (r.get("amount_cents") or 0)
    for j, (k, v) in enumerate(sorted(sums.items()), start + 1):
        ws.cell(row=j, column=2, value=k)
        c = ws.cell(row=j, column=7, value=v / 100)
        c.number_format = "0.00"
    total_row = start + 1 + len(sums) + 1
    ws.cell(row=total_row, column=1, value="合计").font = bold
    t = ws.cell(row=total_row, column=7, value="=SUM(G2:G%d)" % (n + 1))
    t.font = bold
    t.number_format = "0.00"
    widths = {1: 6, 2: 18, 3: 30, 4: 22, 5: 12, 6: 26, 7: 12, 8: 14, 9: 12, 10: 30, 11: 14, 12: 10}
    for k, v in widths.items():
        ws.column_dimensions[ws.cell(row=1, column=k).column_letter].width = v
    # 退回清单
    rejected = [r for r in rows if r.get("reject_reason")]
    ws2 = wb.create_sheet("退回清单")
    ws2.append(["提交人/批次", "文件名", "发票号码", "金额(元)", "退回原因"])
    for c in ws2[1]:
        c.font = bold
    for r in rejected:
        ws2.append([r.get("folder") or "", r.get("fname") or "", r.get("no") or "",
                    (r.get("amount_cents") or 0) / 100, r.get("reject_reason")])
    if not rejected:
        ws2.append(["（无退回票据）", "", "", "", ""])
    for col, w in zip("ABCDE", (18, 30, 22, 12, 16)):
        ws2.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- 合规校验（财务核验环节）----------
def validate_fields(f, today=None):
    """字段自校验。抓不到就不校验，宁可少报也不误报。"""
    issues = []
    no = f.get("no")
    if no and len(no) not in (8, 20):
        issues.append("发票号码位数异常（%d 位，常见为 8 位或数电票 20 位）" % len(no))
    date = f.get("date")
    if date:
        today = today or time.strftime("%Y-%m-%d")
        if date > today:
            issues.append("开票日期晚于今天（%s）" % date)
        elif date < "2010-01-01":
            issues.append("开票日期过早（%s）" % date)
    cents = f.get("amount_cents")
    if cents is not None:
        if cents <= 0:
            issues.append("金额异常（≤0）")
        elif cents > 1000000000:
            issues.append("金额异常偏大，请核对票面")
    excl, tax = f.get("amount_excl_cents"), f.get("tax_cents")
    if cents is not None and excl is not None and tax is not None:
        if abs(excl + tax - cents) > 100:  # 容差 1 元，规避四舍五入噪声
            issues.append("金额自洽未通过：不含税 %.2f + 税额 %.2f ≠ 价税合计 %.2f" % (
                excl / 100, tax / 100, cents / 100))
    return issues


def check_buyer(f, company_names):
    """抬头校验：未配置本公司名称时不校验，避免无意义告警。"""
    names = [n for n in (company_names or []) if n]
    if not names:
        return "未校验", f.get("buyer")
    buyer = f.get("buyer") or ""
    if not buyer:
        return "待核对", ""
    return ("符合" if any(n in buyer for n in names) else "不符"), buyer


# ---------- OFD（数电票官方格式）----------
def _xml_local(tag):
    return tag.split("}")[-1] if "}" in tag else tag


def read_ofd_text(path):
    """解析 OFD 票面文本：OFD 本质是 zip，票面文字在 Content.xml 的 TextCode 中。

    只读不写、不联网、不需要第三方库。解析失败返回空字符串。
    """
    import zipfile
    import xml.etree.ElementTree as ET
    out = []
    try:
        with zipfile.ZipFile(path) as zf:
            names = set(zf.namelist())
            root_name = "OFD.xml"
            if root_name not in names:
                root_name = next((n for n in names if n.lower().endswith("ofd.xml")), None)
            if not root_name:
                return ""
            doc_root = None
            try:
                root = ET.fromstring(zf.read(root_name))
                for el in root.iter():
                    if _xml_local(el.tag) == "DocRoot" and (el.text or "").strip():
                        doc_root = el.text.strip().lstrip("/")
                        break
            except Exception:
                return ""
            if not doc_root:
                return ""
            base = os.path.dirname(doc_root)
            doc_dir = (base + "/") if base else ""
            page_locs = []
            try:
                doc = ET.fromstring(zf.read(doc_root))
                for el in doc.iter():
                    if _xml_local(el.tag) == "Page":
                        loc = el.attrib.get("BaseLoc") or ""
                        if loc:
                            page_locs.append(loc.lstrip("/"))
            except Exception:
                return ""
            for loc in page_locs:
                full = loc if loc in names else (doc_dir + loc)
                if full not in names:
                    continue
                try:
                    content = ET.fromstring(zf.read(full))
                except Exception:
                    continue
                # 按 TextObject 的 Boundary 纵向排序，保证阅读顺序
                items = []
                for obj in content.iter():
                    if _xml_local(obj.tag) != "TextObject":
                        continue
                    y = 0.0
                    try:
                        b = obj.attrib.get("Boundary") or ""
                        if b:
                            y = float(b.split()[1])
                    except (IndexError, ValueError):
                        pass
                    txt = "".join(t.text or "" for t in obj.iter()
                                  if _xml_local(t.tag) == "TextCode")
                    if txt.strip():
                        items.append((y, txt))
                items.sort(key=lambda x: x[0])
                out.extend(t for _y, t in items)
    except Exception:
        return ""
    return "\n".join(out)


# ---------- 文件与 OCR ----------
def file_md5(path, chunk=1 << 20):
    h = hashlib.md5()
    with open(path, "rb") as fh:
        while True:
            b = fh.read(chunk)
            if not b:
                break
            h.update(b)
    return h.hexdigest()


def zhipu_key():
    """读智谱 API Key：进程环境变量优先，再实时读 Windows 注册表。

    实时读注册表是为了绕开 Windows 经典坑：用户设置"用户级环境变量"后，
    已在运行的 explorer/终端不会刷新环境，之后双击启动的程序全部继承
    旧快照，os.environ 里永远看不到新变量。直接查注册表则立即生效，
    无需注销/重启。
    """
    for name in ("ZHIPUAI_API_KEY", "ZAI_API_KEY"):
        v = os.environ.get(name)
        if v and v.strip():
            return v.strip()
    try:
        import winreg
    except ImportError:  # 非 Windows
        return ""
    for root, path in ((winreg.HKEY_CURRENT_USER, "Environment"),
                       (winreg.HKEY_LOCAL_MACHINE,
                        r"SYSTEM\CurrentControlSet\Control\Session Manager\Environment")):
        try:
            with winreg.OpenKey(root, path) as k:
                for name in ("ZHIPUAI_API_KEY", "ZAI_API_KEY"):
                    try:
                        v, _t = winreg.QueryValueEx(k, name)
                    except OSError:
                        continue
                    if v and v.strip():
                        return v.strip().strip('"')
        except OSError:
            continue
    return ""


def mask_key(key):
    """脱敏展示：只露头尾各 4 位。"""
    key = (key or "").strip()
    return key[:4] + "…" + key[-4:] if len(key) >= 10 else ""


def save_zhipu_key(key):
    """把 Key 写入 Windows 当前用户环境变量（等效于旧版 PowerShell 命令，但无需命令行）。

    写注册表 HKCU\\Environment 后广播 WM_SETTINGCHANGE，之后新开的程序都能读到；
    同时更新当前进程 os.environ，本服务立即生效，无需重启。
    """
    key = (key or "").strip()
    if not key:
        raise ValueError("API Key 不能为空")
    try:
        import winreg
        with winreg.CreateKey(winreg.HKEY_CURRENT_USER, "Environment") as k:
            winreg.SetValueEx(k, "ZHIPUAI_API_KEY", 0, winreg.REG_SZ, key)
        try:  # 通知系统刷新环境变量，让之后新启动的程序立刻可见
            import ctypes
            ctypes.windll.user32.SendMessageTimeoutW(
                0xFFFF, 0x001A, 0, "Environment", 0x0002, 1000,
                ctypes.byref(ctypes.c_ulong()))
        except Exception:
            pass
    except ImportError:  # 非 Windows（开发环境）：仅写入当前进程
        pass
    os.environ["ZHIPUAI_API_KEY"] = key


def clear_zhipu_key():
    """删除本机保存的 Key（注册表 + 当前进程）。"""
    try:
        import winreg
        with winreg.CreateKey(winreg.HKEY_CURRENT_USER, "Environment") as k:
            try:
                winreg.DeleteValue(k, "ZHIPUAI_API_KEY")
            except OSError:
                pass
    except ImportError:
        pass
    os.environ.pop("ZHIPUAI_API_KEY", None)


def friendly_ocr_error(error):
    """把供应商错误转换为用户可执行的提示，不暴露响应正文。"""
    code = getattr(error, "code", None)
    if code in (401, 403):
        return "API Key 无效、已删除或无模型权限"
    if code == 429:
        return "调用频率或账户额度受限，请稍后重试"
    text = str(error).lower()
    if "timed out" in text or "timeout" in text:
        return "连接超时，请检查网络后重试"
    return "智谱服务暂不可用，请检查网络或稍后重试"


def zhipu_chat(messages, model="glm-4v-flash", timeout=60, max_tokens=None):
    key = zhipu_key()
    if not key:
        raise RuntimeError("未配置 ZHIPUAI_API_KEY")
    payload = {"model": model, "messages": messages}
    if max_tokens:
        payload["max_tokens"] = max_tokens
    req = urllib.request.Request(
        "https://open.bigmodel.cn/api/paas/v4/chat/completions",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json", "Authorization": "Bearer " + key})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        data = json.loads(resp.read().decode("utf-8"))
    return data["choices"][0]["message"]["content"] or ""


def ocr_b64(b64, mime, model="glm-4v-flash", timeout=60):
    """把 base64 图片发到智谱视觉模型，返回识别出的纯文本。"""
    return zhipu_chat([{
            "role": "user",
            "content": [
                {"type": "image_url", "image_url": {"url": "data:%s;base64,%s" % (mime, b64)}},
                {"type": "text",
                 "text": "请完整识别这张发票/票据上的全部文字，按原文输出纯文本。务必保留「发票号码："
                         "xxx」「开票日期：xxxx年x月x日」「价税合计（小写）¥xx.xx」「销售方名称」「购买方名称」"
                         "「项目名称」下的明细等关键内容。不要解释、不要总结。"},
            ],
        }], model, timeout)


def pick_folder(initial=""):
    """调用系统目录选择器；取消时返回空字符串。"""
    import tkinter as tk
    from tkinter import filedialog
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    try:
        start = initial if initial and os.path.isdir(initial) else os.path.expanduser("~/Desktop")
        return filedialog.askdirectory(initialdir=start, mustexist=True) or ""
    finally:
        root.destroy()


def ocr_image_file(path, model="glm-4v-flash"):
    ext = os.path.splitext(path)[1].lower()
    mime = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png",
            "bmp": "image/bmp", "webp": "image/webp"}.get(ext.lstrip("."), "image/jpeg")
    with open(path, "rb") as fh:
        b64 = base64.b64encode(fh.read()).decode()
    return ocr_b64(b64, mime, model)


# ---------- 本地 OCR（RapidOCR，可选依赖）----------
_RAPID = None
_RAPID_FAILED = False


def rapid_engine():
    """懒加载本地识别引擎；未安装或初始化失败返回 None。"""
    global _RAPID, _RAPID_FAILED
    if RapidOCR is None or _RAPID_FAILED:
        return None
    if _RAPID is None:
        try:
            _RAPID = RapidOCR()
        except Exception:
            _RAPID_FAILED = True
            return None
    return _RAPID


def local_ocr_ready():
    return RapidOCR is not None and not _RAPID_FAILED


# 本机识别组件自动安装：首次启动或升级后若未安装，后台静默装上，用户无感知
OCR_INSTALL = {"status": "idle"}  # idle / installing / ready / failed
_OCR_INSTALLStarted = False


def _boot_mtime():
    """记录启动时 app.py 的修改时间，用于判断「代码已更新但服务未重启」。"""
    try:
        return os.path.getmtime(os.path.abspath(__file__))
    except OSError:
        return 0.0


BOOT_MTIME = _boot_mtime()


def code_stale():
    """app.py 在进程启动之后被改过 → 页面提示用户重启，避免白跑旧代码。"""
    try:
        return abs(os.path.getmtime(os.path.abspath(__file__)) - BOOT_MTIME) > 0.5
    except OSError:
        return False


def _pip_install_ocr():
    import subprocess
    pkg = "rapidocr-onnxruntime>=1.4,<2"
    attempts = [
        [sys.executable, "-m", "pip", "install", pkg,
         "-i", "https://pypi.tuna.tsinghua.edu.cn/simple"],
        [sys.executable, "-m", "pip", "install", pkg],
    ]
    for cmd in attempts:
        try:
            r = subprocess.run(cmd, capture_output=True, timeout=900)
            if r.returncode == 0:
                break
        except Exception:
            continue
    global _RAPID, _RAPID_FAILED
    _RAPID = None       # 重新懒加载，让新装的组件立即生效
    _RAPID_FAILED = False
    OCR_INSTALL["status"] = "ready" if local_ocr_ready() else "failed"


def ensure_local_ocr_async():
    """启动时调用：未安装则后台静默安装，绝不阻断启动与使用。"""
    global _OCR_INSTALLStarted
    if _OCR_INSTALLStarted:
        return
    _OCR_INSTALLStarted = True
    if local_ocr_ready():
        OCR_INSTALL["status"] = "ready"
        return
    def _worker():
        OCR_INSTALL["status"] = "installing"
        _pip_install_ocr()
    threading.Thread(target=_worker, daemon=True).start()


def local_ocr_image(pil_img):
    """对 PIL 图片做本地识别，返回按阅读顺序排列的文本。失败抛异常由调用方兜底。"""
    import numpy as np
    eng = rapid_engine()
    if eng is None:
        raise RuntimeError("本地识别组件未安装")
    img = pil_img.convert("RGB")
    # 长边降采样：显著提速并降低超大图失败率
    long_side = 1600
    w, h = img.size
    if max(w, h) > long_side:
        ratio = long_side / float(max(w, h))
        img = img.resize((int(w * ratio), int(h * ratio)))
    result, _ = eng(np.array(img))
    if not result:
        return ""
    ordered = sorted(result, key=lambda x: (min(p[1] for p in x[0]), min(p[0] for p in x[0])))
    return "\n".join(str(t).strip() for _box, t, _score in ordered if t)


def local_ocr_file(path):
    """图片发票的本地识别。"""
    if Image is None:
        raise RuntimeError("缺少 Pillow 组件")
    with Image.open(path) as im:
        return local_ocr_image(im)


def local_ocr_pdf(path, scale=2.2, max_pages=4):
    """无文本层 PDF：渲染成图后本地识别。"""
    if pdfium is None or Image is None:
        raise RuntimeError("缺少 PDF 渲染组件")
    texts = []
    doc = pdfium.PdfDocument(path)
    try:
        for i in range(min(len(doc), max_pages)):
            texts.append(local_ocr_image(doc[i].render(scale=scale).to_pil()))
    finally:
        doc.close()
    return "\n".join(texts)


def render_pdf_pages(path, scale=2.2, max_pages=4):
    """无文本层 PDF -> 渲染成 PNG -> base64 列表（pypdfium2）。"""
    if pdfium is None:
        return []
    doc = pdfium.PdfDocument(path)
    out = []
    try:
        for i in range(min(len(doc), max_pages)):
            page = doc[i]
            bmp = page.render(scale=scale)
            img = bmp.to_pil()
            buf = io.BytesIO()
            img.convert("RGB").save(buf, format="PNG")
            out.append(base64.b64encode(buf.getvalue()).decode())
    finally:
        doc.close()
    return out


def parse_file(path, ocr_enabled=True, ocr_model="glm-4v-flash", company_names=None):
    """解析单张发票 -> (fields, warn)。warn 非空表示有降级/提示。"""
    ext = os.path.splitext(path)[1].lower()
    warn = None
    text = ""
    pdf_text = ""
    did_ocr = False
    engine_used = None  # local=本机识别 / ai=智谱；用于前端如实标注与缓存自愈
    if ext == ".ofd":
        # 数电票官方格式：直接读 XML 票面文本，最准且不联网、不需要第三方库
        text = read_ofd_text(path)
        if not re.sub(r"\s", "", text or ""):
            warn = "OFD 未解析到票面文本（可能是扫描件版式）"
            if ocr_enabled:
                text = ""  # 交由下方统一兜底逻辑处理
    elif ext == ".pdf":
        if PdfReader is not None:
            try:
                text = "".join((p.extract_text() or "") for p in PdfReader(path).pages)
                pdf_text = text
            except Exception as e:
                warn = "PDF解析异常:%s" % e
        # 文本过少或字符编码损坏时尝试渲染 OCR。
        needs_ocr = len(re.sub(r"\s", "", text or "")) < 40 or text.count("\x00") >= 8
        if needs_ocr and ocr_enabled:
            # 优先本地识别（免 Key、数据不出本机），不可用时再走 AI
            if local_ocr_ready():
                try:
                    local_text = local_ocr_pdf(path)
                    if local_text.strip():
                        text, warn, did_ocr = local_text, None, True
                        engine_used = "local"
                except Exception as e:
                    warn = "本地识别失败：%s" % e
            if not did_ocr and zhipu_key():
                try:
                    pages = render_pdf_pages(path)
                    chunks = []
                    for b64 in pages:
                        try:
                            chunks.append(ocr_b64(b64, "image/png", ocr_model))
                        except Exception as e:
                            warn = "PDF页OCR失败：%s" % friendly_ocr_error(e)
                    if chunks:
                        text = "\n".join(chunks)
                        warn = None
                        did_ocr = True
                        engine_used = "ai"
                except Exception as e:
                    warn = "PDF渲染失败:%s" % e
        if not warn and needs_ocr and not did_ocr:
            warn = "PDF无可用文本层且未完成OCR(需配置ZHIPUAI_API_KEY)"
    else:  # 图片
        if ocr_enabled:
            if local_ocr_ready():
                try:
                    local_text = local_ocr_file(path)
                    if local_text.strip():
                        text, warn, did_ocr = local_text, None, True
                        engine_used = "local"
                except Exception as e:
                    warn = "本地识别失败：%s" % e
            if not did_ocr and zhipu_key():
                try:
                    text = ocr_image_file(path, ocr_model)
                    did_ocr = True
                    engine_used = "ai"
                    warn = None if text.strip() else "OCR返回为空"
                except Exception as e:
                    warn = "OCR失败：%s" % friendly_ocr_error(e)
            if not did_ocr and not warn:
                warn = "图片发票未启用OCR(需配置ZHIPUAI_API_KEY)"
    fields = extract_fields(text or "", os.path.basename(path), company_names)
    if did_ocr and pdf_text:
        fallback = extract_fields(pdf_text, os.path.basename(path), company_names)
        # 数字字段优先采用本地文本层；OCR 主要补齐损坏的中文购销方。
        for key in ("no", "code", "date", "amount_cents"):
            if fallback.get(key) is not None:
                fields[key] = fallback[key]
        for key in ("seller", "buyer"):
            if fields.get(key) is None and fallback.get(key) is not None:
                fields[key] = fallback[key]
    missing = [name for key, name in (("no", "号码"), ("amount_cents", "金额"),
                                      ("date", "日期"), ("seller", "销售方"))
               if fields.get(key) is None]
    if missing:
        detail = "需复核：%s未识别" % "、".join(missing)
        warn = "%s；%s" % (warn, detail) if warn else detail
    fields["class_text"] = (text or "")[:6000]
    fields["ocr"] = did_ocr
    fields["ocr_engine"] = engine_used
    return fields, warn


def fingerprint(fields, md5hex):
    # 永久库条目没有文件内容，md5 为空时不能生成 exact 键，否则会互相误判重复
    keys = [("exact", "md5|%s" % md5hex)] if md5hex else []
    if fields.get("no") and fields.get("amount_cents"):
        keys.append(("high", "no+amt|%s|%s" % (fields["no"], fields["amount_cents"])))
    if fields.get("no"):
        if fields.get("code"):
            keys.append(("high", "code+no|%s|%s" % (fields["code"], fields["no"])))
        else:
            keys.append(("mid", "no|%s" % fields["no"]))
    if fields.get("seller") and fields.get("amount_cents") and fields.get("date"):
        keys.append(("low", "sel+amt+date|%s|%s|%s" % (
            _strip_ws(fields["seller"]).lower(), fields["amount_cents"], fields["date"])))
    return keys


# ---------- 台账 ----------
def _rotate_backup(path, tag=""):
    """写入前滚动备份，保留最近 BACKUP_KEEP 份。

    tag 用于区分来源（如 prerestore），避免同名同秒覆盖掉另一份有效备份。
    """
    if not os.path.exists(path):
        return None
    try:
        os.makedirs(BACKUP_DIR, exist_ok=True)
        name = os.path.basename(path)
        stamp = time.strftime("%Y%m%d-%H%M%S") + ("-" + tag if tag else "")
        dst = os.path.join(BACKUP_DIR, "%s.%s.bak" % (name, stamp))
        i = 1
        while os.path.exists(dst):  # 同秒冲突时绝不覆盖已有备份
            dst = os.path.join(BACKUP_DIR, "%s.%s-%d.bak" % (name, stamp, i))
            i += 1
        shutil.copy2(path, dst)
        olds = sorted(glob.glob(os.path.join(BACKUP_DIR, name + ".*.bak")), reverse=True)
        for old in olds[BACKUP_KEEP:]:
            try:
                os.remove(old)
            except OSError:
                pass
        return dst
    except OSError:
        return None


def list_backups():
    if not os.path.isdir(BACKUP_DIR):
        return []
    items = []
    for p in sorted(glob.glob(os.path.join(BACKUP_DIR, "*.bak")), reverse=True):
        try:
            items.append({"name": os.path.basename(p), "size": os.path.getsize(p),
                          "mtime": time.strftime("%Y-%m-%d %H:%M:%S",
                                                 time.localtime(os.path.getmtime(p)))})
        except OSError:
            continue
    return items[:20]


def load_ledger():
    """读取台账。损坏时保留现场并登记状态，绝不静默当成空台账。"""
    if os.path.exists(LEDGER):
        try:
            with open(LEDGER, encoding="utf-8") as fh:
                data = json.load(fh)
            if isinstance(data, dict) and isinstance(data.get("records"), dict):
                LEDGER_STATUS["error"] = None
                LEDGER_STATUS["record_count"] = len(data["records"])
                return data
            raise ValueError("台账结构异常")
        except Exception as e:
            # 保住损坏现场，避免后续写入把唯一副本覆盖掉
            try:
                os.makedirs(BACKUP_DIR, exist_ok=True)
                shutil.copy2(LEDGER, os.path.join(
                    BACKUP_DIR, "invoice_ledger.corrupt-%s.bak" % time.strftime("%Y%m%d-%H%M%S")))
            except OSError:
                pass
            LEDGER_STATUS["error"] = "台账文件损坏或格式异常（%s），已自动留档，可在「台账恢复」中还原。" % e
            LEDGER_STATUS["backups"] = list_backups()
            return {"version": ENGINE_VER, "records": {}}
    LEDGER_STATUS["error"] = None
    return {"version": ENGINE_VER, "records": {}}


def save_ledger(led):
    _rotate_backup(LEDGER)
    tmp = LEDGER + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(led, fh, ensure_ascii=False, indent=1)
    os.replace(tmp, LEDGER)
    LEDGER_STATUS["record_count"] = len(led.get("records", {}))


def restore_ledger(name):
    """从备份还原台账；只允许还原 backups 目录内的文件。"""
    path = os.path.join(BACKUP_DIR, os.path.basename(name))
    if not os.path.isfile(path) or os.path.dirname(os.path.abspath(path)) != os.path.abspath(BACKUP_DIR):
        raise ValueError("备份文件不存在")
    try:
        data = json.load(open(path, encoding="utf-8"))
        if not (isinstance(data, dict) and isinstance(data.get("records"), dict)):
            raise ValueError("该备份不是有效的台账文件")
    except Exception as e:
        raise ValueError("备份无法解析：%s" % e)
    # 用 prerestore 标签另存当前文件，避免同秒覆盖掉正要还原的那份备份
    _rotate_backup(LEDGER, "prerestore")
    shutil.copy2(path, LEDGER)
    LEDGER_STATUS["error"] = None
    return len(data["records"])


def load_used_ledger():
    """已报销永久库：结构 {version, items:[...]}，脱离文件位置存在。"""
    if os.path.exists(USED_LEDGER):
        try:
            with open(USED_LEDGER, encoding="utf-8") as fh:
                data = json.load(fh)
            if isinstance(data, dict) and isinstance(data.get("items"), list):
                return data
        except Exception:
            pass
    return {"version": 1, "items": []}


def save_used_ledger(data):
    _rotate_backup(USED_LEDGER)
    tmp = USED_LEDGER + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=1)
    os.replace(tmp, USED_LEDGER)


def used_ledger_add(records, batch=""):
    """把一批记录并入已报销永久库，按 (号码+金额) 去重，返回新增条数。"""
    data = load_used_ledger()
    seen = set()
    for it in data["items"]:
        seen.add(used_item_key(it))
    added = 0
    now = time.time()
    for r in records:
        key = used_item_key(r)
        if not key or key in seen:
            continue
        seen.add(key)
        data["items"].append({
            "no": r.get("no"), "code": r.get("code"),
            "amount_cents": r.get("amount_cents"), "date": r.get("date"),
            "seller": r.get("seller"), "buyer": r.get("buyer"),
            "kind": r.get("kind"), "cat_label": r.get("cat_label"),
            "fname": r.get("fname"), "path": r.get("path"),
            "batch": batch or time.strftime("%Y-%m"), "closed_at": now,
        })
        added += 1
    if added:
        save_used_ledger(data)
    return added


def used_item_key(r):
    """永久库唯一键：优先 号码+金额，其次 代码+号码，避免 20 位兜底误抓重复入库。"""
    no = (r or {}).get("no")
    cents = (r or {}).get("amount_cents")
    if no and cents is not None:
        return "no+amt|%s|%s" % (no, cents)
    code = (r or {}).get("code")
    if no and code:
        return "code+no|%s|%s" % (code, no)
    return ""


def load_config():
    if os.path.exists(CONFIG):
        try:
            return json.load(open(CONFIG, encoding="utf-8"))
        except Exception:
            pass
    return {"watch_dirs": [], "used_dirs": [], "archive_dir": "",
            "ocr_model": "glm-4v-flash", "company_names": BUYER_DEFAULT}


def save_config(cfg):
    tmp = CONFIG + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(cfg, fh, ensure_ascii=False, indent=1)
    os.replace(tmp, CONFIG)


def is_within(path, folder):
    try:
        return os.path.commonpath((os.path.abspath(path), os.path.abspath(folder))) == os.path.abspath(folder)
    except (ValueError, OSError):
        return False


def records_in_dirs(led, dirs):
    return [r for r in led.get("records", {}).values()
            if any(is_within(r["path"], d) for d in dirs)]


def collect_records():
    """当前配置目录下的全部台账记录，并按配置标注 is_used。"""
    cfg = load_config()
    led = load_ledger()
    watch_dirs = list(cfg.get("watch_dirs", []))
    used_set = {os.path.normpath(d) for d in cfg.get("used_dirs", []) if d and os.path.isdir(d)}
    dirs = watch_dirs + list(cfg.get("used_dirs", []))
    recs = records_in_dirs(led, dirs)
    for r in recs:
        folder_used = any(is_within(r["path"], d) for d in used_set)
        r["in_watch"] = any(is_within(r["path"], d) for d in watch_dirs)
        r["is_used"] = bool(r.get("used_override")) if r.get("used_override_set") else folder_used
    return recs


def validate_categories(raw):
    if not isinstance(raw, list) or not 1 <= len(raw) <= 50:
        raise ValueError("分类数量应为 1-50 个")
    out, seen = [], set()
    for i, cat in enumerate(raw):
        cid = str(cat.get("id") or "").strip()
        label = str(cat.get("label") or "").strip()
        if not re.fullmatch(r"[a-z0-9_-]{1,40}", cid) or cid in seen or not label:
            raise ValueError("第 %d 个分类的标识或名称无效" % (i + 1))
        kws = [str(x).strip() for x in cat.get("keywords", []) if str(x).strip()]
        out.append({"id": cid, "label": label[:50],
                    "note": str(cat.get("note") or "").strip()[:120],
                    "keywords": list(dict.fromkeys(kws))[:100]})
        seen.add(cid)
    if "other" not in seen:
        raise ValueError("必须保留 other（其他/待分类）")
    return out


def iter_invoice_files(dirs):
    seen = set()
    for d in dirs or []:
        if not d or not os.path.isdir(d):
            continue
        for root, _subs, files in os.walk(d):
            if any(seg.startswith(".") for seg in root.split(os.sep)):
                continue
            for fn in files:
                if os.path.splitext(fn)[1].lower() in INVOICE_EXTS:
                    p = os.path.normpath(os.path.join(root, fn))
                    if p not in seen:
                        seen.add(p)
                        yield p


def build_scan(watch_dirs, used_dirs, ocr_enabled, ocr_model):
    """扫描并增量更新台账（只读文件内容，绝不删除/移动）。"""
    led = load_ledger()
    if led.get("version") != ENGINE_VER:
        led = {"version": ENGINE_VER, "records": {}}
    all_dirs = list(dict.fromkeys(list(watch_dirs or []) + list(used_dirs or [])))
    used_set = {os.path.normpath(d) for d in (used_dirs or []) if d and os.path.isdir(d)}
    if os.path.exists(CONFIG):
        led_cfg = json.load(open(CONFIG, encoding="utf-8"))
    else:
        led_cfg = {}
    buyers = led_cfg.get("company_names") or BUYER_DEFAULT
    categories = get_categories(led_cfg)
    parse_sig = hashlib.sha256(json.dumps({
        "engine": ENGINE_VER, "ocr": bool(ocr_enabled), "ocr_model": ocr_model,
        "ocr_key": bool(zhipu_key()), "buyers": buyers,
    }, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:16]
    cat_sig = hashlib.sha256(json.dumps(categories, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:16]

    # 1) 增量解析（AI OCR 走网络，线程池并发提速；结果按序写回台账）
    todo = []
    for p in iter_invoice_files(all_dirs):
        st = os.stat(p)
        rec = led["records"].get(p)
        unchanged = rec and rec.get("mtime") == st.st_mtime and rec.get("size") == st.st_size
        if unchanged and rec.get("parse_sig") == parse_sig:
            # 一次性故障自愈：上次走了 AI 兜底仍缺关键字段（如销售方），而本机识别现已可用
            # → 本轮强制重解析一次（本机识别免费、离线）。重解析后 engine 变为 local，
            #   即使仍缺字段也不会再进此分支，避免每轮空转。
            _warn_txt = rec.get("warn") or ""
            if ("需复核" in _warn_txt and local_ocr_ready()
                    and (rec.get("ocr_engine") == "ai"
                         or (rec.get("ocr") and rec.get("ocr_engine") is None))):
                todo.append((p, st))
                continue
            if (rec.get("ocr") and rec.get("warn") == "PDF无可用文本层且未完成OCR(需配置ZHIPUAI_API_KEY)"
                    and all(rec.get(k) is not None for k in ("no", "amount_cents", "date", "seller"))):
                rec["warn"] = None
            if rec.get("cat_rule") == "人工":
                rec["cat_label"] = next((c["label"] for c in categories
                                          if c["id"] == rec.get("cat_id")), "其他/待分类")
            elif rec.get("cat_sig") != cat_sig:
                cid, rule = classify_text(rec.get("cls_text", ""), categories)
                rec.update(cat_id=cid, cat_rule=rule,
                           cat_label=next((c["label"] for c in categories if c["id"] == cid), cid),
                           cat_sig=cat_sig)
            continue
        todo.append((p, st))
    if todo:
        from concurrent.futures import ThreadPoolExecutor

        def _parse_one(path):
            try:
                return parse_file(path, ocr_enabled, ocr_model, buyers)
            except Exception as e:  # 文件被占用/删除等，不让单张失败拖垮整轮扫描
                return {"ocr": False}, "解析失败:%s" % e

        with ThreadPoolExecutor(max_workers=4) as ex:
            parsed = list(ex.map(_parse_one, [p for p, _st in todo]))
        for (p, st), (fields, warn) in zip(todo, parsed):
            old = led["records"].get(p)
            md5hex = file_md5(p)
            issues = validate_fields(fields)
            buyer_status, _buyer = check_buyer(fields, buyers)
            fp = fingerprint(fields, md5hex)
            cls_text = os.path.basename(p) + "|" + (fields.get("seller") or "") + "|" + \
                " ".join(fields.get("items") or []) + "|" + \
                " ".join(fields.get("svc") or []) + "|" + (fields.get("kind") or "") + "|" + \
                fields.get("class_text", "")
            rec = {
                "path": p, "folder": os.path.basename(os.path.dirname(p)),
                "fname": os.path.basename(p), "size": st.st_size, "mtime": st.st_mtime,
                "md5": md5hex,
                "no": fields.get("no"), "code": fields.get("code"),
                "date": fields.get("date"), "amount_cents": fields.get("amount_cents"),
                "seller": fields.get("seller"), "buyer": fields.get("buyer"),
                "kind": fields.get("kind"), "items": fields.get("items"),
                "svc": fields.get("svc"),
                "ocr": fields.get("ocr", False), "ocr_engine": fields.get("ocr_engine"),
                "warn": warn,
                "fp": fp, "cls_text": cls_text[:6000], "parse_sig": parse_sig, "cat_sig": cat_sig,
                "cat_id": old.get("cat_id") if old and old.get("cat_rule") == "人工" else None,
                "cat_label": old.get("cat_label") if old and old.get("cat_rule") == "人工" else None,
                "cat_rule": old.get("cat_rule") if old and old.get("cat_rule") == "人工" else None,
                "used_override": old.get("used_override", False) if old else False,
                "used_override_set": old.get("used_override_set", False) if old else False,
                "check_issues": issues, "buyer_status": buyer_status,
                "verify_status": old.get("verify_status", "未查验") if old else "未查验",
                "reject_reason": old.get("reject_reason", "") if old else "",
                "closed": old.get("closed", False) if old else False,
                "file_missing": False,
                "first_seen": old.get("first_seen", time.time()) if old else time.time(),
                "updated": time.time(),
            }
            cid, rule = classify_text(rec["cls_text"], categories)
            if rec["cat_id"] is None:
                rec["cat_id"], rec["cat_rule"] = cid, rule
                rec["cat_label"] = next((c["label"] for c in categories if c["id"] == cid), cid)
            led["records"][p] = rec

    # 2) 文件消失时的处理：已报销/已结案/已查验的记录永久保留（财务留痕），其余才清理
    cfg_prefixes = tuple(os.path.normpath(d) for d in all_dirs if d)
    if cfg_prefixes:
        for p in list(led["records"]):
            if any(is_within(p, d) for d in cfg_prefixes) and not os.path.exists(p):
                r = led["records"][p]
                if r.get("used_override") or r.get("closed") or r.get("verify_status"):
                    r["file_missing"] = True
                else:
                    del led["records"][p]

    # 3) 内容查重（同源不判重：同一物理文件只有一条记录）
    recs = records_in_dirs(led, all_dirs)
    for r in recs:
        folder_used = any(is_within(r["path"], d) for d in used_set)
        r["in_watch"] = any(is_within(r["path"], d) for d in (watch_dirs or []))
        r["in_used_dir"] = folder_used
        r["is_used"] = bool(r.get("used_override")) if r.get("used_override_set") else folder_used
    idx = {}
    for r in recs:
        for w, k in r["fp"]:
            idx.setdefault(k, []).append(r)
    # 已报销永久库注入：原文件即使被移走/删除，仍能持续查重
    for it in load_used_ledger()["items"]:
        ghost = {"path": "used://" + (used_item_key(it) or "unknown"),
                 "origin_path": it.get("path") or "",
                 "folder": "已报销库", "fname": it.get("fname") or "（历史台账记录）",
                 "is_used": True, "ghost": True, "no": it.get("no"),
                 "amount_cents": it.get("amount_cents"), "batch": it.get("batch")}
        for w, k in fingerprint(it, None):
            idx.setdefault(k, []).append(ghost)
    for r in recs:
        r["dups"] = []
        seenp = set()
        for w, k in r["fp"]:
            basis = ("文件内容完全一致" if k.startswith("md5|") else
                     "发票代码+号码一致" if k.startswith("code+no|") else
                     "发票号+金额一致" if k.startswith("no+amt|") else
                     "发票号码一致" if k.startswith("no|") else
                     "销方+金额+日期一致" if k.startswith("sel+amt+date|") else w)
            for o in idx.get(k, []):
                if o["path"] != r["path"] and o["path"] not in seenp:
                    # 永久库里就是它自己时不判重，避免结案后自匹配
                    if o.get("origin_path") and o["origin_path"] == r["path"]:
                        continue
                    seenp.add(o["path"])
                    r["dups"].append({"level": "高危" if w != "low" else "疑似", "basis": basis,
                                      "fname": o["fname"], "folder": o["folder"],
                                      "path": o["path"], "no": o.get("no"),
                                      "amount": o.get("amount_cents"), "is_used": o["is_used"],
                                      "ghost": bool(o.get("ghost")), "batch": o.get("batch", "")})
    save_ledger(led)
    return led, used_set, recs


# ---------- CSV ----------
def csv_escape(v):
    v = "" if v is None else str(v)
    if any(ch in v for ch in (",", '"', "\n")):
        v = '"' + v.replace('"', '""') + '"'
    return v


def export_csv(recs, scope):
    head = ["文件名", "所在文件夹", "金额(元)", "开票日期", "发票号码", "发票代码",
            "销售方", "费用分类", "分类依据", "状态", "重复提示", "OCR/提示", "文件路径"]
    lines = [",".join(head)]
    for r in recs:
        used = "已使用" if r["is_used"] else "待使用"
        dup = "；".join("%s[%s/%s]" % (d["basis"], d["folder"], d["fname"])
                        for d in r["dups"]) or ""
        if (scope == "used" and not r["is_used"]) or \
           (scope == "dup" and not r["dups"]) or \
           (scope == "reused" and (r["is_used"] or not any(d.get("is_used") for d in r["dups"]))) or \
           (scope == "new" and not r.get("in_watch")):
            continue
        row = [r["fname"], r["folder"], (r["amount_cents"] or 0) / 100,
               r.get("date") or "", r.get("no") or "", r.get("code") or "",
               r.get("seller") or "", r.get("cat_label") or "", r.get("cat_rule") or "",
               used, dup, r.get("warn") or "", r["path"]]
        lines.append(",".join(csv_escape(x) for x in row))
    return "\ufeff" + "\n".join(lines)


def record_in_scope(r, scope):
    if scope == "used":
        return r["is_used"]
    if scope == "new":
        return bool(r.get("in_watch"))
    if scope == "dup":
        return bool(r.get("dups"))
    if scope == "reused":
        return (not r["is_used"]) and any(d.get("is_used") for d in (r.get("dups") or []))
    return True


def build_summary_workbook(recs):
    """分类汇总表：分类做列、金额按开票日期竖排、底部合计行，另附明细页。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    ordered = [c["label"] for c in get_categories()]
    groups = {}
    for r in recs:
        cents = r.get("amount_cents")
        if not cents:
            continue
        lbl = r.get("cat_label") or "其他/待分类"
        if lbl not in ordered:
            ordered.append(lbl)
        groups.setdefault(lbl, []).append((r.get("date") or "", cents))
    cols = [lbl for lbl in ordered if groups.get(lbl)]
    for lst in groups.values():
        lst.sort(key=lambda x: (x[0] or "9999-99-99", x[1]))

    wb = Workbook()
    ws = wb.active
    ws.title = "分类汇总"
    bold = Font(bold=True)
    n_rows = max((len(v) for v in groups.values()), default=0)
    for j, lbl in enumerate(cols, 1):
        col = get_column_letter(j)
        head = ws.cell(row=1, column=j, value=lbl)
        head.font = bold
        ws.column_dimensions[col].width = max(12, len(lbl) * 2 + 4)
        for i, (_d, cents) in enumerate(groups[lbl], 2):
            cell = ws.cell(row=i, column=j, value=cents / 100)
            cell.number_format = "0.00"
        sum_row = n_rows + 3
        total = ws.cell(row=sum_row, column=j, value="=SUM(%s2:%s%d)" % (col, col, n_rows + 1))
        total.font = bold
        total.number_format = "0.00"
    if cols:
        last = len(cols) + 1
        head = ws.cell(row=1, column=last, value="总计")
        head.font = bold
        grand = ws.cell(row=n_rows + 3, column=last,
                        value="=SUM(A%d:%s%d)" % (n_rows + 3, get_column_letter(len(cols)), n_rows + 3))
        grand.font = bold
        grand.number_format = "0.00"
        ws.column_dimensions[get_column_letter(last)].width = 14

    ws2 = wb.create_sheet("明细")
    head = ["开票日期", "金额(元)", "费用分类", "销售方", "文件名", "所在文件夹", "状态"]
    for j, h in enumerate(head, 1):
        cell = ws2.cell(row=1, column=j, value=h)
        cell.font = bold
    for i, r in enumerate(sorted(recs, key=lambda x: (x.get("date") or "9999-99-99", x.get("fname") or "")), 2):
        vals = [r.get("date") or "", (r.get("amount_cents") or 0) / 100,
                r.get("cat_label") or "", r.get("seller") or "", r.get("fname") or "",
                r.get("folder") or "", "已使用" if r["is_used"] else "待使用"]
        for j, v in enumerate(vals, 1):
            cell = ws2.cell(row=i, column=j, value=v)
            if j == 2:
                cell.number_format = "0.00"
    for j, w in enumerate((12, 11, 13, 32, 36, 20, 9), 1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    return wb


def export_xlsx(recs):
    from io import BytesIO
    buf = BytesIO()
    build_summary_workbook(recs).save(buf)
    return buf.getvalue()


# ---------- HTTP ----------
class Handler(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def _send(self, code, body, ctype="application/json; charset=utf-8"):
        if isinstance(body, str):
            body = body.encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _read_json(self):
        ln = int(self.headers.get("Content-Length", 0))
        return json.loads(self.rfile.read(ln).decode("utf-8")) if ln else {}

    def do_GET(self):
        u = urlparse(self.path)
        if u.path in ("/", "/index.html"):
            if os.path.exists(STATIC):
                html = open(STATIC, encoding="utf-8").read().replace("{{APP_VERSION}}", APP_VERSION)
                return self._send(200, html,
                                  "text/html; charset=utf-8")
            return self._send(404, "缺少 index.html")
        if u.path == "/api/config":
            cfg = load_config()
            cfg["ocr_key_configured"] = bool(zhipu_key())
            cfg["key_masked"] = mask_key(zhipu_key())
            cfg["app_version"] = APP_VERSION
            cfg["code_stale"] = code_stale()
            cfg["local_ocr_ready"] = local_ocr_ready()
            cfg["local_ocr_status"] = OCR_INSTALL.get("status", "idle")
            return self._send(200, json.dumps(cfg))
        if u.path == "/api/subdirs":
            # 列出所选目录的直接子目录，供「添加目录」多选勾选
            q = parse_qs(u.query)
            base = (q.get("path") or [""])[0]
            found = []
            if base and os.path.isdir(base):
                try:
                    for name in sorted(os.listdir(base)):
                        full = os.path.normpath(os.path.join(base, name))
                        if os.path.isdir(full):
                            found.append(full)
                except OSError:
                    pass
            return self._send(200, json.dumps({"base": base, "dirs": found[:200]}, ensure_ascii=False))
        if u.path == "/api/categories":
            return self._send(200, json.dumps(get_categories(), ensure_ascii=False))
        if u.path == "/api/folders":
            desk = os.path.join(os.path.expanduser("~"), "Desktop")
            found = []
            if os.path.isdir(desk):
                for d in sorted(os.listdir(desk)):
                    full = os.path.join(desk, d)
                    if os.path.isdir(full) and (re.match(r"^20\d{4}", d) or "发票" in d
                                                or "报销" in d or "归档" in d):
                        found.append(full)
            return self._send(200, json.dumps({"desktop": desk, "folders": found}))
        if u.path == "/api/file":
            # 发票原文件预览：仅允许已配置目录内的发票文件，防止任意路径读取
            q = parse_qs(u.query)
            path = (q.get("path") or [""])[0]
            cfg = load_config()
            dirs = [d for d in list(cfg.get("watch_dirs", [])) + list(cfg.get("used_dirs", [])) if d]
            ext = os.path.splitext(path)[1].lower()
            if not path or ext not in INVOICE_EXTS or not os.path.isfile(path) \
                    or not any(is_within(path, d) for d in dirs):
                return self._send(404, json.dumps(
                    {"error": "发票文件不存在或不在已配置目录中"}, ensure_ascii=False))
            ctype = "application/pdf" if ext == ".pdf" else {
                ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
                ".bmp": "image/bmp", ".webp": "image/webp"}.get(ext, "application/octet-stream")
            with open(path, "rb") as fh:
                data = fh.read()
            self.send_response(200)
            self.send_header("Content-Type", ctype)
            self.send_header("Cache-Control", "no-store")
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)
            return
        if u.path == "/api/export.csv":
            q = parse_qs(u.query)
            scope = (q.get("scope") or ["all"])[0]
            return self._send(200, export_csv(collect_records(), scope), "text/csv; charset=utf-8")
        if u.path == "/api/export.xlsx":
            from urllib.parse import quote
            q = parse_qs(u.query)
            scope = (q.get("scope") or ["all"])[0]
            recs = [r for r in collect_records() if record_in_scope(r, scope)]
            try:
                data = export_xlsx(recs)
            except ImportError:
                return self._send(500, json.dumps(
                    {"error": "缺少 openpyxl 组件，请先执行: pip install openpyxl"}, ensure_ascii=False))
            self.send_response(200)
            self.send_header("Content-Type",
                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header(
                "Content-Disposition",
                'attachment; filename="invoice_summary.xlsx"; filename*=UTF-8\'\''
                + quote("发票分类汇总.xlsx"))
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)
            return
        if u.path == "/api/ledger-status":
            return self._send(200, json.dumps({
                "error": LEDGER_STATUS.get("error"),
                "record_count": LEDGER_STATUS.get("record_count", 0),
                "used_count": len(load_used_ledger().get("items", [])),
                "backups": list_backups(),
            }, ensure_ascii=False))
        if u.path == "/api/used-ledger":
            items = load_used_ledger().get("items", [])
            return self._send(200, json.dumps(
                {"total": len(items), "items": items[-500:]}, ensure_ascii=False))
        if u.path == "/api/export.worksheet":
            recs = [r for r in collect_records()]
            try:
                data = export_worksheet(recs)
            except ImportError:
                return self._send(500, json.dumps(
                    {"error": "缺少 openpyxl 组件，请先执行: pip install openpyxl"}, ensure_ascii=False))
            from urllib.parse import quote
            self.send_response(200)
            self.send_header("Content-Type",
                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header(
                "Content-Disposition",
                'attachment; filename="invoice_worksheet.xlsx"; filename*=UTF-8\'\''
                + quote("发票审核底稿.xlsx"))
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)
            return
        self._send(404, "{}")

    def do_POST(self):
        u = urlparse(self.path)
        if u.path == "/api/categories":
            try:
                cats = validate_categories(self._read_json().get("categories"))
                with STATE_LOCK:
                    cfg = load_config()
                    cfg["categories"] = cats
                    save_config(cfg)
                return self._send(200, json.dumps({"ok": True, "categories": cats}, ensure_ascii=False))
            except (ValueError, TypeError) as e:
                return self._send(400, json.dumps({"error": str(e)}, ensure_ascii=False))
        if u.path == "/api/scan":
            body = self._read_json()
            watch = body.get("watch_dirs") or []
            used = body.get("used_dirs") or []
            ocr = bool(body.get("ocr_enabled", True))
            model = body.get("ocr_model") or "glm-4v-flash"
            cfg = load_config()
            cfg.update({"watch_dirs": watch, "used_dirs": used,
                        "ocr_model": model, "ocr_enabled": ocr,
                        "archive_dir": body.get("archive_dir") or ""})
            if body.get("company_names"):
                cfg["company_names"] = body["company_names"]
            with STATE_LOCK:
                save_config(cfg)
                led, used_set, recs = build_scan(watch, used, ocr, model)
            catalog = get_categories(cfg)
            cats = {}
            for c in catalog:
                cats[c["id"]] = {"label": c["label"], "note": c["note"], "count": 0, "sum": 0.0}
            n_dup, n_used, n_warn, n_amt = 0, 0, 0, 0
            for r in recs:
                if r["is_used"]:
                    n_used += 1
                if r["dups"]:
                    n_dup += 1
                if r.get("warn"):
                    n_warn += 1
                if r.get("amount_cents") is not None:
                    n_amt += 1
                cid = r.get("cat_id") or "other"
                if cid in cats:
                    cats[cid]["count"] += 1
                    cats[cid]["sum"] += (r["amount_cents"] or 0) / 100
            recs.sort(key=lambda r: (-1 if r["dups"] else 0, -1 if r["is_used"] else 0,
                                     -(r["amount_cents"] or 0)))
            pairs = {tuple(sorted((r["path"], d["path"]))) for r in recs for d in r["dups"]}
            reused = sum(r.get("in_watch") and any(d.get("is_used") for d in r["dups"]) for r in recs)
            watch_count = sum(bool(r.get("in_watch")) for r in recs)
            used_count = sum(bool(r.get("in_used_dir")) for r in recs)
            ai_error = next((r.get("warn") for r in recs if r.get("warn") and
                             ("OCR失败" in r["warn"] or "PDF页OCR失败" in r["warn"])), "")
            return self._send(200, json.dumps({
                "records": recs, "categories": cats,
                "stats": {"total": len(recs), "dup": n_dup, "used": n_used,
                          "warn": sum(bool(r.get("in_watch") and r.get("warn")) for r in recs),
                          "warn_all": n_warn, "amount_ok": n_amt,
                          "ocr_key": bool(zhipu_key()), "dup_pairs": len(pairs),
                          "reused": reused, "watch_count": watch_count,
                          "used_count": used_count, "ai_error": ai_error},
                "config": {**cfg, "ocr_key_configured": bool(zhipu_key())},
            }, ensure_ascii=False))
        if u.path == "/api/pick-folder":
            try:
                path = pick_folder(self._read_json().get("initial") or "")
                return self._send(200, json.dumps({"path": path}, ensure_ascii=False))
            except Exception:
                return self._send(500, json.dumps(
                    {"error": "无法打开 Windows 文件夹选择器，请直接粘贴目录路径"}, ensure_ascii=False))
        if u.path == "/api/ocr-status":
            if not zhipu_key():
                return self._send(200, json.dumps({"ok": False, "configured": False,
                    "message": "尚未配置智谱 API Key"}, ensure_ascii=False))
            body = self._read_json()
            try:
                zhipu_chat([{"role": "user", "content": "仅回复OK"}],
                           body.get("model") or "glm-4v-flash", timeout=20, max_tokens=2)
                return self._send(200, json.dumps({"ok": True, "configured": True,
                    "message": "连接正常"}, ensure_ascii=False))
            except Exception as e:
                return self._send(200, json.dumps({"ok": False, "configured": True,
                    "message": friendly_ocr_error(e)}, ensure_ascii=False))
        if u.path == "/api/save-key":
            body = self._read_json()
            if body.get("clear"):
                clear_zhipu_key()
                return self._send(200, json.dumps(
                    {"ok": True, "configured": False, "masked": "",
                     "message": "已清除本机保存的 API Key"}, ensure_ascii=False))
            key = (body.get("api_key") or "").strip().strip('"').strip("'")
            if len(key) < 10:
                return self._send(400, json.dumps(
                    {"error": "这串内容不像完整的 API Key（应是一长串字母数字），请回网页重新「复制」再粘贴"},
                    ensure_ascii=False))
            save_zhipu_key(key)
            try:
                zhipu_chat([{"role": "user", "content": "仅回复OK"}],
                           "glm-4v-flash", timeout=20, max_tokens=2)
                return self._send(200, json.dumps(
                    {"ok": True, "configured": True, "masked": mask_key(key),
                     "message": "连接正常，AI 识别已就绪"}, ensure_ascii=False))
            except Exception as e:
                return self._send(200, json.dumps(
                    {"ok": False, "configured": True, "masked": mask_key(key),
                     "message": friendly_ocr_error(e)}, ensure_ascii=False))
        if u.path == "/api/override":
            body = self._read_json()
            path = body.get("path")
            with STATE_LOCK:
                led = load_ledger()
                if path not in led["records"]:
                    return self._send(404, json.dumps({"error": "发票记录不存在"}, ensure_ascii=False))
                r = led["records"][path]
                if "cat_id" in body:
                    catalog = get_categories()
                    if body["cat_id"] not in {c["id"] for c in catalog}:
                        return self._send(400, json.dumps({"error": "分类不存在"}, ensure_ascii=False))
                    r["cat_id"] = body["cat_id"]
                    r["cat_label"] = next((c["label"] for c in catalog
                                           if c["id"] == body["cat_id"]), body["cat_id"])
                    r["cat_rule"] = "人工"
                if "used" in body:
                    r["used_override"] = bool(body["used"])
                    r["used_override_set"] = True
                if "verify_status" in body:
                    if body["verify_status"] not in VERIFY_CHOICES:
                        return self._send(400, json.dumps({"error": "查验状态无效"}, ensure_ascii=False))
                    r["verify_status"] = body["verify_status"]
                if "reject_reason" in body:
                    r["reject_reason"] = body["reject_reason"] if body["reject_reason"] in REJECT_CHOICES else ""
                save_ledger(led)
            return self._send(200, json.dumps({"ok": True}))
        if u.path == "/api/review":
            """批量设置查验状态 / 退回原因（财务核验）。"""
            body = self._read_json()
            paths = body.get("paths") or []
            with STATE_LOCK:
                led = load_ledger()
                n = 0
                for p in paths:
                    r = led["records"].get(p)
                    if not r:
                        continue
                    if body.get("verify_status") in VERIFY_CHOICES:
                        r["verify_status"] = body["verify_status"]
                    if "reject_reason" in body:
                        r["reject_reason"] = body["reject_reason"] if body["reject_reason"] in REJECT_CHOICES else ""
                    n += 1
                save_ledger(led)
            return self._send(200, json.dumps({"ok": True, "updated": n}, ensure_ascii=False))
        if u.path == "/api/restore-ledger":
            try:
                n = restore_ledger(self._read_json().get("name") or "")
                return self._send(200, json.dumps(
                    {"ok": True, "records": n, "message": "已还原 %d 条台账记录，请重新查重。" % n},
                    ensure_ascii=False))
            except ValueError as e:
                return self._send(400, json.dumps({"error": str(e)}, ensure_ascii=False))
        if u.path == "/api/import-history":
            """N3 冷启动：导入财务已有的历史台账（CSV/Excel 另存为 CSV）。"""
            body = self._read_json()
            items = parse_history_csv(body.get("csv") or "")
            if not items:
                return self._send(400, json.dumps(
                    {"error": "未识别到有效记录，请确认文件含发票号码或「金额+开票日期」两列"},
                    ensure_ascii=False))
            with STATE_LOCK:
                added = used_ledger_add(items, body.get("batch") or "历史导入")
            return self._send(200, json.dumps(
                {"ok": True, "total": len(items), "added": added,
                 "message": "解析 %d 条，新增 %d 条进入已报销库（重复自动跳过）。" % (len(items), added)},
                ensure_ascii=False))
        if u.path == "/api/close-batch":
            """N1 结案：本批入库已报销 + 可选规范命名归档，一次走完 SOP 最后一环。"""
            import shutil as _shutil
            body = self._read_json()
            paths = body.get("paths") or []
            batch = (body.get("batch") or "").strip() or time.strftime("%Y-%m")
            do_archive = bool(body.get("archive"))
            with STATE_LOCK:
                led = load_ledger()
                picked = [led["records"][p] for p in paths if p in led["records"]]
                if not picked:
                    return self._send(400, json.dumps({"error": "请先勾选要结案的发票"}, ensure_ascii=False))
                for r in picked:
                    r["closed"] = True
                    r["closed_at"] = time.time()
                    r["used_override"] = True
                    r["used_override_set"] = True
                    r["is_used"] = True
                save_ledger(led)
                added = used_ledger_add(picked, batch)
                copied, skipped = [], []
                if do_archive:
                    target = (load_config().get("archive_dir") or "").strip()
                    if not target or not os.path.isdir(target):
                        skipped.append("归档目录未设置或不存在，已跳过归档")
                    else:
                        for r in picked:
                            src = r.get("path")
                            if not src or not os.path.isfile(src):
                                skipped.append(r.get("fname") or src)
                                continue
                            dst = os.path.join(target, archive_name(r))
                            base, ext = os.path.splitext(dst)
                            i = 1
                            while os.path.exists(dst):
                                dst = "%s(%d)%s" % (base, i, ext)
                                i += 1
                            _shutil.copy2(src, dst)
                            copied.append(os.path.basename(dst))
            # 无号码且无金额的票无法生成唯一键，必须明确告知，否则下月查不出来
            no_key = [r.get("fname") or r.get("path") for r in picked if not used_item_key(r)]
            return self._send(200, json.dumps(
                {"ok": True, "closed": len(picked), "added": added,
                 "no_key": no_key, "copied": copied, "skipped": skipped}, ensure_ascii=False))
        if u.path == "/api/extract":
            import shutil
            body = self._read_json()
            paths = body.get("paths") or []
            target = body.get("target_dir") or ""
            if not target or not os.path.isdir(target):
                return self._send(400, json.dumps({"error": "归档目录不存在或未填写"}))
            cfg = load_config()
            allowed = {r["path"] for r in records_in_dirs(
                load_ledger(), list(cfg.get("watch_dirs", [])) + list(cfg.get("used_dirs", [])))}
            copied, skipped = [], []
            for p in paths:
                if p in allowed and os.path.isfile(p):
                    dst = os.path.join(target, os.path.basename(p))
                    if os.path.exists(dst):
                        skipped.append(dst)
                    else:
                        shutil.copy2(p, dst)
                        copied.append(dst)
            return self._send(200, json.dumps({"copied": copied, "skipped": skipped}, ensure_ascii=False))
        self._send(404, "{}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--port", type=int, default=8765)
    ap.add_argument("--no-browser", action="store_true")
    ap.add_argument("--test", nargs="*", help="命令行只读自测目录")
    a = ap.parse_args()
    if a.test is not None:
        cfg = load_config()
        dirs = a.test or cfg.get("watch_dirs") or []
        led, used_set, recs = build_scan(dirs, cfg.get("used_dirs") or [], True,
                                         cfg.get("ocr_model") or "glm-4v-flash")
        for r in sorted(recs, key=lambda r: (r["folder"], r["fname"])):
            dup = ("重复[%s]" % ", ".join("%s/%s:%s" % (d["folder"], d["fname"], d["basis"])
                                          for d in r["dups"])) if r["dups"] else ""
            print("%-6s %-24s %8s %-10s %-20s %-14s %-20s %s%s" % (
                r["folder"], r["fname"][:24],
                "%.2f" % ((r["amount_cents"] or 0) / 100) if r["amount_cents"] is not None else "-",
                r.get("date") or "-", (r.get("no") or "-")[:20],
                (r.get("cat_label") or "")[:14], (r.get("seller") or "-")[:20],
                dup, (" | " + r["warn"]) if r.get("warn") else ""))
        print("\n合计 %d 张 | 重复风险 %d | 已使用 %d | OCR可用=%s | 金额已解析 %d" % (
            len(recs), sum(1 for x in recs if x["dups"]),
            sum(1 for x in recs if x["is_used"]), bool(zhipu_key()),
            sum(1 for x in recs if x.get("amount_cents") is not None)))
        return 0
    srv = ThreadingHTTPServer(("127.0.0.1", a.port), Handler)
    url = "http://127.0.0.1:%d" % a.port
    ensure_local_ocr_async()  # 未装本机识别组件时后台静默安装，用户无感知
    print("发票管家已启动: %s  (Ctrl+C 退出)" % url)
    if not a.no_browser:
        threading.Timer(0.6, lambda: webbrowser.open(url)).start()
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        print("\n已退出")
    return 0


if __name__ == "__main__":
    sys.exit(main())
